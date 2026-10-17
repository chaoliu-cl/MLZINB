from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


ROOT = Path(__file__).resolve().parent
PROFILE = ROOT / "bikeshare_profile"
FIG = PROFILE / "figures"
OUT = ROOT / "Capital_Bikeshare_MLZINB_Contextual_Moderation_Walkthrough.docx"

NAVY = "17324D"
TEAL = "1F7A8C"
BLUE = "2E75B6"
ORANGE = "D97706"
LIGHT_BLUE = "EAF3F7"
LIGHT_TEAL = "E8F4F3"
LIGHT_ORANGE = "FFF3E0"
GRAY = "5B6573"
LIGHT_GRAY = "F3F5F7"
RED = "A33A3A"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_text(cell, text, bold=False, color=None, size=8.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=8.3, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color="FFFFFF", size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i], header_fill)
        if widths:
            hdr.cells[i].width = Inches(widths[i])
    # Repeat the header row in multi-page tables.
    trPr = hdr._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, bold=(first_col_bold and i == 0), size=font_size)
            set_cell_shading(cells[i], "FFFFFF" if ridx % 2 == 0 else LIGHT_GRAY)
            if widths:
                cells[i].width = Inches(widths[i])
        # Keep rows intact when possible.
        trPr = table.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"val": "single", "sz": "4", "color": "D5DCE3"},
                            bottom={"val": "single", "sz": "4", "color": "D5DCE3"},
                            left={"val": "single", "sz": "4", "color": "D5DCE3"},
                            right={"val": "single", "sz": "4", "color": "D5DCE3"})
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Aptos"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    return p


def add_figure(doc, filename, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(width))
    add_caption(doc, caption)


def add_callout(doc, title, text, fill=LIGHT_TEAL, border=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Keep the callout as one visual unit across pages.
    trPr = table.rows[0]._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": "12", "color": border},
                    bottom={"val": "single", "sz": "12", "color": border},
                    left={"val": "single", "sz": "12", "color": border},
                    right={"val": "single", "sz": "12", "color": border})
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.08
    for run in p2.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.06
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(9.4)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.name = "Aptos"
        r.font.size = Pt(9.6)
        rest = text[len(bold_lead):]
        r2 = p.add_run(rest)
        r2.font.name = "Aptos"
        r2.font.size = Pt(9.6)
    else:
        r = p.add_run(text)
        r.font.name = "Aptos"
        r.font.size = Pt(9.6)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def set_doc_defaults(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [("Title", 25, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 11.5, TEAL), ("Heading 3", 10.2, NAVY)]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            # Remove the default Title-style bottom rule; whitespace and typography
            # provide the separation on the cover page.
            ppr = style._element.get_or_add_pPr()
            pbdr = ppr.find(qn("w:pBdr"))
            if pbdr is not None:
                ppr.remove(pbdr)

    # Header and footer.
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("ML-ZINB contextual moderation walkthrough  |  Capital Bikeshare")
    hr.font.name = "Aptos"
    hr.font.size = Pt(7.5)
    hr.font.color.rgb = RGBColor.from_string(GRAY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("UCI Bike Sharing dataset  |  ")
    fr.font.name = "Aptos"
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(footer)
    for r in footer.runs:
        r.font.name = "Aptos"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(GRAY)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("COMPONENT-AWARE EVALUATION")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(8)
    p.add_run("Contextual moderation in a multilevel zero-inflated negative-binomial model")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("A complete walkthrough using the UCI Bike Sharing dataset")
    r.font.name = "Aptos"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    add_callout(doc, "Executive conclusion", "Working-day context changes the within-day clock-time profile of casual bike-rental counts. The result is strongly supported in a 2011 discovery sample, reproduced in a temporally separated 2012 validation sample, and remains supported when uncertainty is made robust to day-level clustering and 14-day calendar dependence. The count component carries most of the evidence; the inflation component also contributes, but its interpretation remains conditional on the fitted zero-inflation specification and on residual distributional checks.", fill=LIGHT_TEAL, border=TEAL)

    add_para(doc, "This report documents the full practical workflow: data and hierarchy audit; sample definition; matched additive, count-only, inflation-only and joint ML-ZINB fits; component-specific tests; response-scale predictions; adequacy diagnostics; sensitivity analysis; influence checks; and a reproducible interpretation of the moderation result.")

    meta = add_table(doc, ["Item", "Description"], [
        ["Dataset", "UCI Bike Sharing, classified by UCI as Social Science"],
        ["Outcome", "Hourly casual-user rental counts"],
        ["Level 1 / Level 2", "Recorded clock-hour / calendar day"],
        ["Primary sample", "15,720 hours in 655 complete 24-hour days"],
        ["Contextual moderator", "Working day versus nonworking day"],
        ["Report date", "16 October 2026"],
    ], widths=[1.55, 4.8], header_fill=TEAL, font_size=8.7, first_col_bold=True)
    add_para(doc, "Prepared from the verified analysis bundle in 09_MODERATION_DATASET_SEARCH. The report is an applied methodological walkthrough, not a claim of causal transport effects or a general population estimate.")
    doc.add_page_break()


def build_report():
    doc = Document()
    set_doc_defaults(doc)
    add_title_page(doc)

    add_heading(doc, "How to read this walkthrough", 1)
    add_para(doc, "The central estimand is a contextual moderation of a within-day profile. The working-day indicator is constant within a calendar day, while the clock-time basis varies across the 24 hourly observations in that day. The interaction therefore asks whether the entire daily profile changes with context. The ML-ZINB mixture is evaluated in its two modeled components: the negative-binomial count mean and the extra-zero probability.")
    add_bullet(doc, "Primary inference uses complete days so every cluster contributes the same 24 hourly occasions.")
    add_bullet(doc, "Discovery uses 2011; temporal validation uses 2012; the full sample is shown for precision and descriptive completeness.")
    add_bullet(doc, "The interaction is tested as a vector of clock-time products, not as a single scalar slope.")
    add_bullet(doc, "The latent inflation probability is a model quantity. It is not treated as a verified class of people who never rent bikes.")

    add_heading(doc, "1. Dataset, hierarchy and research question", 1)
    add_para(doc, "The UCI Bike Sharing dataset contains hourly Capital Bikeshare records from 2011 and 2012, augmented with calendar and weather variables. The outcome used here is casual-user rental count, so the units are rental events rather than unique riders. The substantive question is: does working-day context modify the daily time pattern of casual rentals?")
    add_para(doc, "The hierarchy is direct and inspectable: hourly observations are nested within calendar days. The contextual moderator is workingday, defined by the source as a day that is neither a weekend nor a holiday. It varies across days and is constant within each day. Clock-time terms vary within every day, making the interaction estimable without conflating a between-day mean with a within-day profile.")
    add_table(doc, ["Sample", "Hours", "Complete days", "Zeros", "Mean", "Variance"], [
        ["2011 discovery", "7,320", "305", "590 (8.06%)", "31.780", "1,644.787"],
        ["2012 validation", "8,400", "350", "688 (8.19%)", "43.649", "3,308.540"],
        ["Full matched sample", "15,720", "655", "1,278 (8.13%)", "38.122", "2,568.714"],
    ], widths=[1.45, 0.9, 1.15, 1.2, 0.9, 1.05], header_fill=TEAL, font_size=8.6, first_col_bold=True)
    add_para(doc, "The raw download contains 17,379 hourly records across 731 days. Seventy-six days are incomplete because the source is short of a nominal 24-hour grid by 165 records. The primary sample retains only the 655 complete days; an all-released-days sensitivity keeps the incomplete days without silently turning missing hours into zeros. Hourly totals reconcile exactly to the companion daily file for casual, registered and total rentals.")
    add_figure(doc, "structure_and_observed_patterns.png", "Figure 1. Grouping support and observed daily profiles. The selected sample gives each primary cluster exactly 24 hourly records and retains the observed zero counts.")

    add_heading(doc, "2. Search record and frozen analysis boundary", 1)
    add_para(doc, "The dataset was selected after searching for a social or health-science example that could demonstrate the practical contextual-moderation toolkit and that would reproduce the interaction in a later sample. The selection record is part of the result: failed candidates and rejected specifications are retained rather than hidden.")
    add_table(doc, ["Candidate or specification", "Finding", "Decision"], [
        ["German health-care panel", "Discovery joint interaction p = .07066", "Rejected as the supported demonstration"],
        ["Bikeshare midday-window model", "Count evidence was strong, but the inflation interaction had SE about 10,000", "Rejected as the regular four-model example"],
        ["Bikeshare smooth daily profile", "Strong discovery evidence and temporal replication in 2012", "Selected for the walkthrough"],
    ], widths=[1.75, 3.1, 1.65], header_fill=ORANGE, font_size=8.4, first_col_bold=True)
    add_para(doc, "The selected daily-profile equations and complete-day sample rule were frozen before inspecting the 2012 interaction results. The temporal check is therefore a useful validation boundary, although the overall exercise remains exploratory model development rather than a preregistered study. A simple three-attempt Bonferroni threshold would be .0167; the primary and calendar-HAC joint tests are far below that threshold.")
    add_callout(doc, "Why the specification was revised", "The original midday window placed the inflation product in a region where modeled extra zeros were essentially absent. The resulting standard error of about 10,000 signaled weak identification. The daily-profile amendment tests the entire clock-time shape and lets the data identify where the component changes occur.", fill=LIGHT_ORANGE, border=ORANGE)

    add_heading(doc, "3. Model specification", 1)
    add_para(doc, "The four matched models differ only in which contextual clock-time products are included. All models use the same observations, controls, random-intercept structure and dispersion parameterization within a sample. This makes the nested comparisons interpretable as tests of the contextual products.")
    add_para(doc, "For hour i in day j, the count component is written as:")
    add_callout(doc, "Count component", "log(mu_ij) = beta_0 + u_j + f(t_i) + beta_W Working_j + delta' [f(t_i) x Working_j] + controls_ij,   u_j ~ Normal(0, sigma_u^2)", fill=LIGHT_BLUE, border=BLUE)
    add_para(doc, "The inflation component is:")
    add_callout(doc, "Inflation component", "logit(pi_ij) = gamma_0 + g(t_i) + gamma_W Working_j + kappa' [g(t_i) x Working_j] + controls_ij.   Y_ij follows a ZINB2 mixture with mean mu_ij, zero-inflation probability pi_ij and dispersion phi.", fill=LIGHT_BLUE, border=BLUE)
    add_para(doc, "Here f(t) contains the first two sine/cosine harmonics and g(t) contains the first harmonic. These bases have exactly zero mean within a complete 24-record day, so the products describe a within-day profile rather than a hidden between-day mean. The controls are month indicators, temperature and its square, humidity, wind speed and bad-weather status. Full-sample equations also include a 2012 indicator.")
    add_table(doc, ["Model", "Count contextual products", "Inflation contextual products"], [
        ["Additive", "None", "None"],
        ["Count only", "4 clock-time x workingday products", "None"],
        ["Inflation only", "None", "2 clock-time x workingday products"],
        ["Joint", "4 products", "2 products"],
    ], widths=[1.15, 2.75, 2.6], header_fill=TEAL, font_size=8.5, first_col_bold=True)

    add_heading(doc, "4. Matched model fits and numerical checks", 1)
    add_para(doc, "All four primary fits converged with positive-definite Hessians. The joint model has the lowest AIC among the nested ML-ZINBs. A same-count-specification NB2 comparator has AIC 112,833.13 versus 111,662.85 for the joint ML-ZINB, which supports the mixture in this working specification. This comparison is descriptive model selection evidence; no ordinary boundary chi-square test is claimed for NB versus ZINB.")
    add_table(doc, ["Model", "logLik", "AIC", "df", "Converged", "pdHess", "Max |gradient|"], [
        ["Additive", "-57,061.740", "114,189.479", "33", "Yes", "Yes", "0.089"],
        ["Count only", "-55,844.487", "111,762.974", "37", "Yes", "Yes", "0.045"],
        ["Inflation only", "-56,984.511", "114,039.022", "35", "Yes", "Yes", "0.093"],
        ["Joint", "-55,792.427", "111,662.854", "39", "Yes", "Yes", "0.016"],
        ["NB2 comparator", "-56,387.566", "112,833.132", "29", "Yes", "Yes", "0.150"],
    ], widths=[1.2, 1.2, 1.25, 0.55, 0.8, 0.65, 1.15], header_fill=NAVY, font_size=8.1, first_col_bold=True)
    add_para(doc, "The day marginal-likelihood reconstruction is a key audit step. Group scores include fixed effects, dispersion and the random-intercept standard deviation. The reconstructed likelihood sum matches glmmTMB within 4.22e-10; summed day scores match the package gradient within 1.94e-05. These checks support the subsequent day-cluster sandwich and calendar-HAC covariance calculations.")

    add_heading(doc, "5. Moderation evidence: discovery, validation and component location", 1)
    add_para(doc, "The primary test compares the additive model with the joint model. It has 6 degrees of freedom because the contextual profile is a vector: four count products plus two inflation products. The result is overwhelming in both years and in the full matched sample.")
    add_table(doc, ["Sample", "Additive -> joint LR", "df", "p"], [
        ["2011 discovery", "1,349.297", "6", "2.30e-288"],
        ["2012 validation", "1,290.666", "6", "1.14e-275"],
        ["Full matched sample", "2,538.625", "6", "<1e-300"],
    ], widths=[1.75, 1.7, 0.65, 2.1], header_fill=TEAL, font_size=8.6, first_col_bold=True)
    add_para(doc, "Conditional additions locate the evidence while holding the other component's products present. The count additions are strongly supported, and the inflation additions are also supported, although the count component has a more stable and substantively direct interpretation.")
    add_table(doc, ["Sample", "Comparison", "LR", "df", "p"], [
        ["2011", "Count -> joint", "43.416", "2", "3.74e-10"],
        ["2011", "Inflation -> joint", "1,285.534", "4", "4.56e-277"],
        ["2012", "Count -> joint", "55.155", "2", "1.05e-12"],
        ["2012", "Inflation -> joint", "1,213.108", "4", "2.29e-261"],
        ["Full", "Count -> joint", "104.120", "2", "2.46e-23"],
        ["Full", "Inflation -> joint", "2,384.168", "4", "<1e-300"],
    ], widths=[0.8, 1.75, 1.0, 0.5, 2.0], header_fill=NAVY, font_size=8.2, first_col_bold=True)
    add_para(doc, "Because observations within a day and adjacent calendar days can be dependent, the robust confirmation uses day-cluster sandwich covariance and a 14-day Bartlett calendar-HAC covariance. The calendar-HAC retains dates' actual calendar positions, including gaps from excluded days.")
    add_table(doc, ["Sample", "Component", "Wald", "df", "p (14-day calendar HAC)"], [
        ["2011 discovery", "Joint", "790.219", "6", "2.00e-167"],
        ["2011 discovery", "Count", "713.906", "4", "3.40e-153"],
        ["2011 discovery", "Inflation", "28.498", "2", "6.48e-07"],
        ["2012 validation", "Joint", "567.151", "6", "2.83e-119"],
        ["2012 validation", "Count", "440.440", "4", "5.06e-94"],
        ["2012 validation", "Inflation", "26.190", "2", "2.06e-06"],
        ["Full", "Joint", "1,076.459", "6", "2.58e-229"],
        ["Full", "Count", "945.816", "4", "1.97e-203"],
        ["Full", "Inflation", "53.260", "2", "2.72e-12"],
    ], widths=[1.35, 1.0, 1.0, 0.5, 2.55], header_fill=ORANGE, font_size=7.9, first_col_bold=True)
    add_callout(doc, "Inference statement", "The 2012 validation preserves both the contextual count-profile interaction and the modeled inflation-profile interaction after allowing for day-level clustering and 14-day calendar dependence. The result is therefore much stronger than a single model-based p-value, while still remaining an association in one transport system.", fill=LIGHT_TEAL, border=TEAL)

    add_heading(doc, "6. Component-aware predictions and visual interpretation", 1)
    add_para(doc, "A practical ML-ZINB walkthrough should move from link-scale tests to quantities that an applied researcher can see. For each reference record, the analysis computes the count-process mean mu, the modeled structural-zero probability pi, the overall expectation (1 - pi) mu, and the total zero probability pi + (1 - pi)[phi/(phi + mu)]^phi. The mixture is formed before averaging across records.")
    add_para(doc, "The profiles below average each context over the same empirical distribution of calendar and weather covariates. Day random effects are fixed at zero. Bands are pointwise 95% calendar-HAC intervals conditional on the observed reference covariates. These are common-reference model associations, not causal counterfactuals and not predictions integrated over the random-effect distribution.")
    add_figure(doc, "component_predictions.png", "Figure 2. Count mean, modeled extra-zero probability, observable expectation and total-zero probability by clock hour and working-day context.")
    add_table(doc, ["Hour", "Context", "Expected rentals", "Total-zero probability"], [
        ["03:00", "Nonworking", "6.468", "0.068"],
        ["08:00", "Nonworking", "21.965", "0.006"],
        ["13:00", "Nonworking", "155.307", "0.000007"],
        ["18:00", "Nonworking", "67.455", "0.000234"],
        ["22:00", "Nonworking", "33.018", "0.002"],
        ["03:00", "Working", "1.787", "0.513"],
        ["08:00", "Working", "18.950", "0.009"],
        ["13:00", "Working", "39.374", "0.001"],
        ["18:00", "Working", "43.697", "0.001"],
        ["22:00", "Working", "18.471", "0.009"],
    ], widths=[0.8, 1.25, 2.0, 2.15], header_fill=TEAL, font_size=8.2, first_col_bold=False)
    add_para(doc, "The response-scale contrast is especially clear around daytime hours. The count mean is much higher for nonworking days around 13:00, whereas the working-day profile shifts mass toward the morning and evening commuting periods. The total-zero probability is high in the working-day early-morning profile and very small around midday for both contexts. A latent pi curve must not be read as a measured proportion of people who never cycle.")
    add_figure(doc, "context_contrasts.png", "Figure 3. Working-day minus nonworking-day contrasts for the joint model, with additive-model overlays. The joint interaction tests establish moderation on the link scale; the plot shows its response-scale magnitude.")
    add_para(doc, "A descriptive window summary comparing hours 10-15 with the other 18 hours gives a difference-in-contrasts of -72.064 rentals for the count mean, -0.055 for pi, -71.954 for the observable expectation and -0.071 for total-zero probability. This is a communication aid, not the primary harmonic-product test.")

    add_heading(doc, "7. Model adequacy and what it limits", 1)
    add_para(doc, "The primary model was checked with 500 unconditional simulations, drawing fresh day effects while holding fitted parameters fixed. With 15,720 observations, residual tests are sensitive to small discrepancies, so the statistics should be read together with the envelopes and the hour/day pattern figures.")
    add_table(doc, ["Diagnostic", "Statistic", "p"], [
        ["Uniformity", "0.023", "1.80e-07"],
        ["Dispersion", "0.511", "<.004 (500 simulations)"],
        ["Zero inflation", "1.066", "0.008"],
        ["Day-aggregated uniformity", "0.056", "0.035"],
    ], widths=[2.4, 1.3, 2.65], header_fill=ORANGE, font_size=8.5, first_col_bold=True)
    add_table(doc, ["Statistic", "Observed", "Simulation median", "95% simulation interval"], [
        ["Mean", "38.122", "38.421", "37.240 - 39.648"],
        ["Variance", "2,568.714", "3,431.289", "3,085.225 - 3,926.933"],
        ["Zero fraction", "0.081", "0.076", "0.073 - 0.080"],
        ["Positive mean", "41.496", "41.594", "40.332 - 42.951"],
        ["Maximum", "367", "907", "693.850 - 1,390.675"],
    ], widths=[1.65, 1.25, 1.7, 1.75], header_fill=NAVY, font_size=8.2, first_col_bold=True)
    add_figure(doc, "calibration_primary.png", "Figure 4. Primary-model calibration: observed counts, zeros and positive counts against unconditional simulation envelopes.")
    add_figure(doc, "residual_patterns_primary.png", "Figure 5. Primary-model residual patterns by clock hour and calendar dependence.")
    add_para(doc, "The most material discrepancy is the upper tail: the observed maximum is 367, while the primary-model simulation median maximum is about 907 and the lower envelope is about 694. The model also shows dispersion and residual-uniformity problems. These findings constrain claims about the full distribution and about probabilities in the extreme tail. They do not erase the repeated contextual profile interaction, but they mean the mixture component curves should be presented as conditional model summaries rather than fully validated probabilities.")

    add_heading(doc, "8. Sensitivity analysis and influential days", 1)
    add_para(doc, "The sensitivity suite addresses alternative within-day heterogeneity, random structures, flexible clock-time shapes, weather-by-context differences, incomplete-day coverage and optimization. Each valid variant includes the contextual products appropriate to that model. AIC values are compared only within the same observation set.")
    add_table(doc, ["Variant", "Rows", "AIC", "pdHess", "Max |gradient|", "Context test"], [
        ["Random count slope", "15,720", "110,042.176", "Yes", "0.048", "p < 1e-300"],
        ["Flexible hour joint", "15,720", "108,421.285", "Yes", "0.014", "p < 1e-300"],
        ["Third count harmonic", "15,720", "109,516.993", "Yes", "0.061", "p < 1e-300"],
        ["Second inflation harmonic", "15,720", "111,646.013", "Yes", "0.013", "p < 1e-300"],
        ["Temperature x context", "15,720", "111,639.591", "Yes", "0.098", "p < 1e-300"],
        ["All released days", "17,379", "120,611.225", "Yes", "0.070", "p < 1e-300"],
        ["BFGS optimizer", "15,720", "111,662.854", "Yes", "0.020", "p < 1e-300"],
        ["Varying dispersion", "15,720", "109,898.898", "Yes", "0.070", "p < 1e-300"],
    ], widths=[1.7, 0.75, 1.2, 0.7, 1.1, 1.05], header_fill=TEAL, font_size=7.7, first_col_bold=True)
    add_figure(doc, "sensitivity_tests.png", "Figure 6. Contextual profile tests across model variants. The variants differ in test dimension, so the p-values are evidence of persistence, not a single common effect size.")
    add_para(doc, "The varying-dispersion model improves zero calibration (zero-inflation diagnostic p about .884) but still has upper-tail and residual-uniformity issues. An added inflation random intercept collapses near zero and has a non-positive-definite Hessian, so it is excluded from inference. The flexible-hour additive comparator has a weak inflation nuisance direction with a maximum SE around 567; its conventional LR reference is not used as confirmation. These are reasons to keep the component-specific primary inference transparent.")
    add_figure(doc, "calibration_varying_dispersion.png", "Figure 7. Varying-dispersion sensitivity: zero calibration improves, but upper-tail calibration remains imperfect.")

    add_heading(doc, "Influence and deletion checks", 2)
    add_para(doc, "First-order deletion influence was evaluated for all 655 days. Five largest approximate influences and five fixed calendar positions were then refit exactly, giving 10 distinct deletion refits. The largest selected exact shift is 0.380 model standard errors; the approximation/exact standardized-shift correlation is about 0.9994. No selected deletion removes the joint moderation evidence. This is not a claim of 655 exact deletion refits or of a full refit cluster bootstrap.")
    add_table(doc, ["Date", "Selection", "Max shift (model SE)", "Converged", "pdHess", "Joint p"], [
        ["2011-05-02", "Top approximate influence", "0.380", "Yes", "Yes", "<1e-300"],
        ["2012-10-28", "Top approximate influence", "0.317", "Yes", "Yes", "<1e-300"],
        ["2012-04-22", "Top approximate influence", "0.294", "Yes", "Yes", "<1e-300"],
        ["2011-07-30", "Top approximate influence", "0.266", "Yes", "Yes", "<1e-300"],
        ["2012-09-15", "Top approximate influence", "0.239", "Yes", "Yes", "<1e-300"],
        ["2011-01-01", "Fixed calendar position", "0.120", "Yes", "Yes", "<1e-300"],
        ["2012-07-15", "Fixed calendar position", "0.192", "Yes", "Yes", "<1e-300"],
    ], widths=[1.1, 2.0, 1.3, 0.75, 0.7, 1.0], header_fill=ORANGE, font_size=7.9, first_col_bold=True)
    add_figure(doc, "influential_days.png", "Figure 8. Largest first-order day-deletion influences and exact refit checks.")

    add_heading(doc, "9. Practical toolkit for applied researchers", 1)
    add_para(doc, "The Bikeshare analysis illustrates a short, reproducible sequence for checking contextual moderation in any ML-ZINB application:")
    add_bullet(doc, "Audit the hierarchy: verify that the contextual variable is constant within clusters, that the focal predictor varies within clusters, and that cluster sizes and missingness support the intended contrast.")
    add_bullet(doc, "Fit four matched models: additive, count-only interaction, inflation-only interaction and joint interaction. Use the same rows and controls within each comparison.")
    add_bullet(doc, "Test the interaction vector jointly, then locate the evidence by component. In a harmonic profile, report the whole profile rather than interpreting one sine or cosine coefficient as a morning effect.")
    add_bullet(doc, "Plot count mean, inflation probability, observable expectation and total-zero probability on a common reference distribution of covariates. Include an additive overlay so response-scale nonlinearity is visible.")
    add_bullet(doc, "Use group-aware uncertainty. Reconstruct cluster scores when possible, then report a cluster sandwich and a calendar-HAC or other time-aware covariance when clusters are ordered.")
    add_bullet(doc, "Check simulation envelopes, residual patterns, dispersion and zero calibration before giving substantive meaning to the mixture components.")
    add_bullet(doc, "Probe alternative random structures, profile bases, missing-cluster rules and optimization. Treat convergence and Hessian flags as identification evidence, not as housekeeping.")
    add_bullet(doc, "Run influence checks at the cluster level and state exactly how many deletions were approximated versus refit.")
    add_callout(doc, "Recommended applied conclusion", "The clock-time profile of casual bike rentals differs between working and nonworking days. This contextual association is reproduced in a subsequent year and persists under dependence-aware uncertainty. Component-aware analysis distinguishes the count-profile change from the modeled extra-zero change; conclusions about the latter remain conditional on the inflation specification and adequacy checks.", fill=LIGHT_TEAL, border=TEAL)

    add_heading(doc, "10. Limitations and reproducibility", 1)
    add_para(doc, "This is an association in one transport system. It does not establish causal moderation, a verified structural-zero mechanism, or generalization to other cities. The context is day-level rather than neighborhood- or institution-level, so the example is most appropriate when the methodological study allows repeated occasions within calendar or organizational units.")
    add_para(doc, "Other limitations include changing system size and seasonality, unavailable inventory or rider exposure, missing clock-hour records, possible daylight-saving aggregation effects, dependence beyond the chosen HAC lags and the assumptions of the ZINB mixture. The complete-day primary sample changes the population of days; the all-released-days sensitivity is therefore material.")
    add_para(doc, "Source and citation: UCI Bike Sharing dataset, DOI 10.24432/C5W894, classified by UCI as Social Science. Cite Fanaee-T and Gama (2013), Event labeling combining ensemble detectors and background knowledge, DOI 10.1007/s13748-013-0040-3. UCI states a CC BY 4.0 license. The original analysis archive, download timestamps, URLs and SHA-256 hashes are retained in the accompanying bundle.")
    add_para(doc, "Reproduction materials: the analysis protocol, adequacy amendment, raw outputs, figures, source manifest and bundle verification are in 09_MODERATION_DATASET_SEARCH/bikeshare_profile and 09_MODERATION_DATASET_SEARCH/outputs. The verification record reports all checks passed, including source hashes, data-structure checks, matched-model tests, score reconstruction, prediction identities, exact deletion coverage, sensitivity coverage, diagnostics and embedded report links.")

    add_heading(doc, "Appendix: key numerical takeaways", 1)
    add_table(doc, ["Quantity", "Value", "Interpretation"], [
        ["Discovery LR test", "p = 2.30e-288", "Strong joint moderation in 2011"],
        ["Validation LR test", "p = 1.14e-275", "Strong joint moderation in 2012"],
        ["Validation HAC joint test", "p = 2.83e-119", "Persists with day/calendar dependence"],
        ["Validation HAC count test", "p = 5.06e-94", "Count profile is strongly moderated"],
        ["Validation HAC inflation test", "p = 2.06e-06", "Inflation profile also contributes"],
        ["Observed zero fraction", "8.13%", "Enough observed zeros for a mixture walkthrough"],
        ["Observed maximum", "367", "Upper tail remains difficult for the primary model"],
        ["Primary simulated median maximum", "907", "Substantial upper-tail misfit"],
        ["Largest selected deletion shift", "0.380 model SEs", "No selected day drives the result"],
    ], widths=[2.0, 1.55, 3.0], header_fill=NAVY, font_size=8.2, first_col_bold=True)

    # Core properties and metadata.
    props = doc.core_properties
    props.title = "Contextual moderation in a multilevel zero-inflated negative-binomial model: Capital Bikeshare walkthrough"
    props.subject = "Component-aware ML-ZINB methodological walkthrough"
    props.author = ""
    props.keywords = "ML-ZINB, contextual moderation, multilevel, bike sharing, zero inflation"
    props.comments = "Generated from the verified analysis bundle."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
