from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.enum.section import WD_ORIENT
from pathlib import Path

out = Path(r'C:\Users\chaoliu\Downloads\MLZINB_Publication_Bundle_ADD_MEPS_CCD_Simulation_2026-08-12\04_COMBINED_WRITING\final_manuscript\Model_Specification_Contextual_Moderation_MLZINB.docx')
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

# Theme and style helpers
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08

for name, size, color, space_before, space_after in [
    ('Title', 20, '000000', 0, 8),
    ('Heading 1', 14, '17365D', 14, 6),
    ('Heading 2', 11.5, '17365D', 10, 4),
    ('Heading 3', 10.5, '3D556E', 8, 3),
]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(space_before)
    st.paragraph_format.space_after = Pt(space_after)
    st.paragraph_format.keep_with_next = True

# Custom styles
if 'Subtitle Custom' not in styles:
    st = styles.add_style('Subtitle Custom', WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    st.font.size = Pt(11)
    st.font.italic = True
    st.font.color.rgb = RGBColor(89, 89, 89)
    st.paragraph_format.space_after = Pt(14)
if 'Equation Block' not in styles:
    st = styles.add_style('Equation Block', WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Cambria Math'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.35)
    st.paragraph_format.right_indent = Inches(0.25)
    st.paragraph_format.space_before = Pt(5)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.0
if 'Table Text' not in styles:
    st = styles.add_style('Table Text', WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    st.font.size = Pt(8.6)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
if 'Table Header' not in styles:
    st = styles.add_style('Table Header', WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    st.font.size = Pt(8.6)
    st.font.bold = True
    st.font.color.rgb = RGBColor(255,255,255)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

# Helpers
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)
    shd.set(qn('w:val'), 'clear')

def set_cell_border(cell, color='D9D9D9', sz='4'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_row_cant_split(row):
    """Keep a table row intact when a table crosses a page boundary."""
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    trPr.append(cant_split)

def set_col_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(int(width*1440)))
            tcW.set(qn('w:type'), 'dxa')

def add_body(text, bold_lead=None):
    p = doc.add_paragraph(style='Normal')
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def add_eq(text):
    p = doc.add_paragraph(style='Equation Block')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Cambria Math'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
    return p

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style='Normal')
        p.style = styles['Normal']
        p.paragraph_format.left_indent = Inches(0.28 + 0.22*level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        p.add_run('• ' + item)

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, '17365D')
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = styles['Table Header']
        p.add_run(h)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, val in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, 'F3F6F9' if ridx % 2 == 1 else 'FFFFFF')
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = styles['Table Text']
            p.add_run(str(val))
    set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89,89,89)
    return p

# Footer with page numbers
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('Model specification of contextual moderation in ML-ZINB  •  ')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(127,127,127)
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

# Title block
p = doc.add_paragraph(style='Title')
p.add_run('Model Specification of Contextual Moderation in Multilevel Zero Inflated Negative Binomial Models')
p = doc.add_paragraph(style='Subtitle Custom')
p.add_run('Manuscript ready methods section')

add_body('This section specifies how contextual moderation is represented, estimated, tested, and interpreted in a multilevel zero-inflated negative-binomial (ML-ZINB) model. The specification separates the conditional count process from the structural-zero process and defines contextual moderation as an interaction between a within-cluster predictor and a cluster-level moderator. This separation is essential because a predictor can alter the expected count among observations in the count-generating state, the probability of structural-zero membership, or both. The two components therefore require separate estimands, diagnostics, tests, and substantive interpretations.')

add_body('The notation below is written for a two-level design, with observations i nested in clusters j. It extends directly to additional levels by adding random effects and indices to the relevant component. The model can be used for individuals within schools, patients within providers, districts within states, or repeated observations within higher-level units.')

doc.add_heading('1 Purpose and estimands', level=1)
add_body('The outcome Yᵢⱼ is a nonnegative count with an excess of zeros relative to a conventional negative-binomial model. The ML-ZINB model treats each zero as arising from one of two latent states: a structural-zero state, in which the outcome is deterministically zero, or a count-generating state, in which zero is still possible because the negative-binomial distribution has positive probability at zero. This distinction is a modeling device and should be given a substantive interpretation only when the study design supports a meaningful separation between the two states.')
add_body('Let Xᵢⱼ be the focal lower-level predictor and Zⱼ the contextual moderator. The primary moderation estimands are the coefficient of the within-cluster-by-context product in the count component and the coefficient of the corresponding product in the inflation component. These coefficients answer different questions. The count-component interaction describes how context changes the association between the within-cluster predictor and the conditional count intensity. The inflation-component interaction describes how context changes the association between the within-cluster predictor and the log odds of structural-zero membership.')

add_table(
    ['Symbol', 'Definition', 'Role in the model'],
    [
        ('Yᵢⱼ', 'Observed nonnegative count for observation i in cluster j', 'Outcome'),
        ('Xᵢⱼ', 'Focal lower-level predictor', 'Individual or unit-level exposure'),
        ('Zⱼ', 'Cluster-level contextual moderator', 'Context that changes the focal association'),
        ('Wᵢⱼ', 'Within-cluster deviation Xᵢⱼ − X̄ⱼ', 'Primary within-cluster predictor'),
        ('Bⱼ', 'Cluster mean X̄ⱼ', 'Between-cluster predictor retained for decomposition'),
        ('μᵢⱼ', 'Conditional mean in the count-generating state', 'Count-component target'),
        ('πᵢⱼ', 'Probability of structural-zero membership', 'Inflation-component target'),
        ('uᶜⱼ, uᶻⱼ', 'Random intercepts for count and inflation links', 'Unobserved cluster heterogeneity'),
        ('eᵢⱼ', 'Exposure or person-time', 'Count-component offset, when applicable'),
    ], [0.8, 2.75, 3.45])
add_caption('Table 1. Core notation for contextual moderation in a two-level ML-ZINB model.')

doc.add_heading('2 The zero inflated negative binomial distribution', level=1)
add_body('Conditional on the random effects and covariates, the outcome follows a two-part mixture. With probability πᵢⱼ it belongs to the structural-zero state. With probability 1 − πᵢⱼ it follows a negative-binomial distribution with conditional mean μᵢⱼ and dispersion parameter κ (or an equivalent parameterization used by the software).')
add_eq('P(Yᵢⱼ = y | ·) = πᵢⱼ 1(y = 0) + (1 − πᵢⱼ) NB(y; μᵢⱼ, κ),   y = 0, 1, 2, …')
add_body('Under the NB2 parameterization, the conditional variance in the count-generating state is μᵢⱼ + μᵢⱼ²/κ. The resulting observed zero probability is not πᵢⱼ alone. It combines structural zeros and sampling zeros from the negative-binomial state:')
add_eq('P(Yᵢⱼ = 0 | ·) = πᵢⱼ + (1 − πᵢⱼ) NB(0; μᵢⱼ, κ)')
add_body('The conditional expected count is likewise the product of the probability of entering the count-generating state and the count-state mean:')
add_eq('E(Yᵢⱼ | ·) = (1 − πᵢⱼ) μᵢⱼ')
add_body('These identities explain why coefficient signs should not be interpreted from the count component alone. A positive coefficient in the inflation component increases the odds of structural-zero membership, whereas a positive coefficient in the count component increases the conditional count intensity. Their effects on the observed mean can therefore point in opposite directions or offset one another.')

doc.add_heading('3 Contextual decomposition of the focal predictor', level=1)
add_body('A contextual moderation analysis should distinguish within-cluster variation in the focal predictor from between-cluster composition. Decompose the predictor as follows:')
add_eq('Xᵢⱼ = Wᵢⱼ + Bⱼ,     Wᵢⱼ = Xᵢⱼ − X̄ⱼ,     Bⱼ = X̄ⱼ')
add_body('The cluster mean Bⱼ should be included alongside Wᵢⱼ in both components when the scientific question concerns the within-cluster association. This adjustment prevents a contextual association from being conflated with the fact that some clusters contain systematically higher or lower values of the focal predictor. If Xᵢⱼ is already a binary indicator, Bⱼ is the cluster proportion exposed to the indicator and Wᵢⱼ is the individual deviation from that proportion.')
add_body('Center the context variable before forming the interaction when a meaningful reference context is available. For a continuous moderator, define Zᶜⱼ = Zⱼ − c, where c may be the sample mean, a prespecified policy value, or another substantively interpretable reference. For a binary moderator, retain a transparent 0/1 coding unless a different contrast is needed. The interaction must be formed from the within-cluster predictor and the centered contextual variable: Wᵢⱼ Zᶜⱼ.')
add_body('The cluster mean may also be allowed to interact with context when the between-cluster association is itself expected to vary across contexts. That term answers a different question from Wᵢⱼ Zᶜⱼ and should not be silently substituted for it. If the purpose is specifically contextual moderation of a within-cluster association, the focal term is Wᵢⱼ Zᶜⱼ and the between-cluster interaction should be omitted or prespecified as a separate estimand.')

add_table(
    ['Specification choice', 'Recommended implementation', 'Interpretive consequence'],
    [
        ('Within predictor', 'Wᵢⱼ = Xᵢⱼ − X̄ⱼ', 'Effect compares units within the same cluster, holding cluster composition fixed'),
        ('Between predictor', 'Include Bⱼ = X̄ⱼ as a main effect', 'Separates contextual composition from the within-cluster association'),
        ('Context scale', 'Center Zⱼ at a meaningful reference', 'Main effect of Wᵢⱼ is the association at the reference context'),
        ('Moderation term', 'Wᵢⱼ × Zᶜⱼ', 'Contextual change in the within-cluster association'),
        ('Binary context', 'Use 0/1 coding with the reference group stated', 'Interaction compares within-cluster slopes across context groups'),
        ('Exposure', 'Use log(eᵢⱼ) as an offset in the count component', 'Models a rate or count per unit exposure'),
    ], [1.45, 3.1, 2.45])
add_caption('Table 2. Recommended construction of contextual moderation terms.')

doc.add_heading('4 Two component link specification', level=1)
doc.add_heading('4.1 Count component', level=2)
add_body('The count component models the expected outcome conditional on membership in the count-generating state. A decomposed contextual moderation specification is:')
add_eq('log(μᵢⱼ) = log(eᵢⱼ) + β₀ + βᵂ Wᵢⱼ + βᴮ Bⱼ + βᶻ Zᶜⱼ + βᵂᶻ Wᵢⱼ Zᶜⱼ + γᵀ Cᵢⱼ + uᶜⱼ')
add_body('Here Cᵢⱼ denotes prespecified individual- and cluster-level covariates. The offset term is included only when the outcome is a count accumulated over unequal exposure or person-time. If no exposure exists, omit the offset rather than inserting a unit exposure merely for notation. The coefficient βᵂ is the within-cluster association at the reference context Zᶜⱼ = 0. The coefficient βᵂᶻ is the change in that log-rate association for a one-unit increase in the centered context.')
add_body('For a one-unit increase in Wᵢⱼ, the conditional rate ratio at context z is exp(βᵂ + βᵂᶻ z). The interaction should therefore be interpreted through conditional slopes or predicted rates at selected contexts, rather than as a stand-alone rate ratio that is assumed to be constant across the context distribution.')

doc.add_heading('4.2 Inflation component', level=2)
add_body('The inflation component models the probability of structural-zero membership:')
add_eq('logit(πᵢⱼ) = α₀ + αᵂ Wᵢⱼ + αᴮ Bⱼ + αᶻ Zᶜⱼ + αᵂᶻ Wᵢⱼ Zᶜⱼ + δᵀ Cᵢⱼ + uᶻⱼ')
add_body('The coefficient αᵂ is the within-cluster association with the log odds of structural-zero membership at the reference context. The coefficient αᵂᶻ is the change in that log-odds association for a one-unit increase in context. The structural-zero odds ratio for a one-unit increase in Wᵢⱼ at context z is exp(αᵂ + αᵂᶻ z). A positive value indicates greater odds of structural-zero membership; it does not indicate greater observed zeros in every situation because the negative-binomial sampling-zero probability also contributes to the observed zero rate.')
add_body('The count and inflation components may share the same fixed-effect covariates, but they do not need to do so. A parsimonious inflation model may be justified when the structural-zero mechanism has a distinct substantive explanation. Any difference in covariate sets should be prespecified and described explicitly, because changing the inflation specification changes the meaning of the latent class and can alter the count-component estimates.')

doc.add_heading('4.3 Random effects and cross component dependence', level=2)
add_body('The standard ML-ZINB specification includes a random intercept in each link. The random effects can be modeled as independent or correlated:')
add_eq('uⱼ = (uᶜⱼ, uᶻⱼ)ᵀ  ~  N(0, Σ),     Σ = [ σ²ᶜ       ρᶜᶻ σᶜ σᶻ ;  ρᶜᶻ σᶜ σᶻ       σ²ᶻ ]')
add_body('The count random intercept captures unobserved cluster differences in conditional count intensity. The inflation random intercept captures unobserved cluster differences in the odds of structural-zero membership. A random inflation intercept should be considered when clusters differ materially in their zero-generating process; forcing inflation heterogeneity to zero can make measured contextual terms absorb unmodeled cluster differences. If the sample or software cannot identify the correlation parameter reliably, fit independent random intercepts and report that restriction.')
add_body('A random slope can be added when the within-cluster association is expected to vary across clusters:')
add_eq('log(μᵢⱼ) = … + (βᵂ + bᵂⱼ) Wᵢⱼ,     bᵂⱼ ~ N(0, σ²ᵂ)')
add_body('A random slope changes the interpretation of the fixed within-cluster coefficient: βᵂ remains the average slope at the reference context, while σ²ᵂ describes unexplained cluster-to-cluster variation around that average. Random slopes should be added because the design or scientific theory supports them, not only after inspecting whether a fixed interaction is significant. They also require sufficient within-cluster variation and enough clusters for stable variance estimation.')

doc.add_heading('5 Model set for evaluating contextual moderation', level=1)
add_body('A component-aware analysis should compare a small, nested set of models that differ only in the focal moderation terms. Let A denote the additive model containing Wᵢⱼ, Bⱼ, Zᶜⱼ, and covariates in the prespecified components. Let C add Wᵢⱼ Zᶜⱼ to the count component, I add it to the inflation component, and J add it to both components.')
add_table(
    ['Model', 'Count component', 'Inflation component', 'Purpose'],
    [
        ('A Additive', 'No W × Z interaction', 'No W × Z interaction', 'Baseline contextual adjustment'),
        ('C Count moderation', 'Add W × Z', 'No W × Z interaction', 'Test count-side moderation alone'),
        ('I Inflation moderation', 'No W × Z interaction', 'Add W × Z', 'Test inflation-side moderation alone'),
        ('J Joint moderation', 'Add W × Z', 'Add W × Z', 'Estimate and evaluate both components together'),
        ('NB comparator', 'Count model without inflation mixture', 'Not applicable', 'Assess whether zero inflation is needed'),
        ('Reduced inflation', 'As prespecified', 'Fixed or random-intercept inflation only', 'Assess sensitivity to inflation heterogeneity'),
    ], [1.35, 2.05, 2.05, 2.55])
add_caption('Table 3. Nested model set for component-aware moderation evaluation.')
add_body('The four focal ZINB models should use identical observations, offsets, control variables, and random-effects structure. This makes likelihood-ratio comparisons interpretable as tests of the added moderation terms rather than tests that combine moderation with a change in sample or baseline model. The NB comparator is useful for assessing the need for the mixture distribution, but an NB-versus-ZINB likelihood-ratio test can involve boundary parameters and should not automatically be treated as an ordinary chi-square test.')

add_body('A practical primary workflow is a gated component procedure. First, compare A and J with a two-degree-of-freedom likelihood-ratio test. If the omnibus test is not supported, report no supported moderation while retaining component estimates and uncertainty. If it is supported, compare J with C to assess the inflation addition and compare J with I to assess the count addition. Apply a multiplicity adjustment such as Holm to these two conditional tests. This produces a transparent distinction between evidence that moderation exists and evidence about which component carries it.')
add_body('The procedure should retain an unlocalized category for the case in which the omnibus test is supported but neither conditional addition is supported after adjustment. That outcome is a limitation of component localization, not evidence that both effects are absent. A numerical failure should be reported separately from a valid nonsignificant test. AIC can be reported as a prediction-oriented model comparison, but it should not be described as a replacement for the prespecified component tests.')

doc.add_heading('6 Estimation and numerical diagnostics', level=1)
add_body('ML-ZINB models are estimated by maximizing the marginal likelihood after integrating over the random effects. In practice, the estimation algorithm and quadrature or approximation settings should be recorded because the likelihood can be difficult when the count and inflation processes are both heterogeneous. The fitted model should be considered usable for inference only after the optimizer, Hessian, covariance, and random-effect diagnostics have been examined.')
add_bullets([
    'Confirm that the optimizer reports convergence and that the estimated Hessian and fixed-effect covariance are finite and positive definite.',
    'Inspect large coefficients, large standard errors, nearly singular covariance matrices, and random-effect standard deviations near zero or at an upper boundary.',
    'Record warnings and failed attempts. If a retry is used, retain the same formula and data; do not silently simplify the model to obtain a successful fit.',
    'Verify that every nested comparison uses the same observations and the same offset definition.',
    'Check observed zero proportions, all-zero clusters, cluster sizes, within-cluster predictor variation, and support of the contextual moderator.',
    'When using cluster-robust or HAC covariance, include nuisance parameters in the score contributions and state the cluster ordering used for any HAC calculation.',
])
add_body('For finite samples with few clusters, model-based standard errors and likelihood-ratio tests may be poorly calibrated. Cluster-robust covariance can provide a useful sensitivity analysis, but it does not repair a misspecified zero-generating mechanism, omitted random slopes, informative cluster size, or a mismatch between the fitted latent classes and the data-generating process. When cluster ordering has substantive meaning, an ordered-cluster HAC comparison may be reported; when clusters are exchangeable, HAC should be described as a calibration comparison rather than as a required correction.')

add_heading = doc.add_heading
add_heading('7 Interpretation on the response scale', level=1)
add_body('Because both components enter the observed mean and zero probability nonlinearly, component coefficients should be complemented by response-scale predictions. For a selected covariate profile, calculate μᵢⱼ and πᵢⱼ from the fitted links and then report the implied conditional expected count and total-zero probability. For a multilevel model, distinguish conditional predictions with random effects set to zero from marginal predictions that integrate over the random-effect distribution. The latter are preferable when the scientific target concerns a new or average cluster.')
add_eq('E(Y | x, z) = E[(1 − π(x, z, uᶻ)) μ(x, z, uᶜ)]')
add_eq('P(Y = 0 | x, z) = E[π(x, z, uᶻ) + (1 − π(x, z, uᶻ)) NB(0; μ(x, z, uᶜ), κ)]')
add_body('When the count and inflation random effects are correlated, the marginal expected count generally cannot be recovered as the product of separately averaged components. The joint distribution of the two random effects must be integrated. This is particularly important for opposite-sign moderation, where the count interaction can increase the count-state intensity while the inflation interaction changes the probability of entering that state in the opposite direction.')
add_body('A useful summary is a contextual difference-in-differences on the response scale. For contexts z₁ and z₀ and focal values w₁ and w₀, define:')
add_eq('DID = [g(w₁, z₁) − g(w₀, z₁)] − [g(w₁, z₀) − g(w₀, z₀)]')
add_body('where g may be the marginal expected count, structural-zero probability, or total-zero probability. The selected values should be reported, and uncertainty can be obtained with the delta method or a simulation-based interval. Response-scale contrasts are summaries of the full model; they do not replace the component-specific link-scale estimands.')

add_heading('8 Recommended reporting language', level=1)
add_body('A complete report should state the centering and decomposition of the focal predictor, the reference value for the context variable, the covariates in each component, the offset, the random-effects structure, and the dispersion parameterization. It should report the count and inflation interaction estimates with standard errors and confidence intervals, followed by component-specific nested tests and the omnibus decision. The report should then translate the findings into predicted means and zero probabilities at substantively meaningful combinations of the focal predictor and context.')
add_body('The following template can be adapted for an applied manuscript:')
add_eq('The focal predictor was decomposed into its within-cluster deviation and cluster mean. Contextual moderation was represented by the within-cluster deviation × contextual moderator product in the count and inflation components. The count component used a log link and the inflation component used a logit link. Both components included the prespecified covariates and cluster random intercepts; the count model included log exposure as an offset when appropriate. We compared additive, count-moderation, inflation-moderation, and joint-moderation ML-ZINB models using the same observations and random-effects structure. An omnibus likelihood-ratio test assessed whether either interaction improved fit, followed by multiplicity-adjusted conditional tests that localized evidence to the count component, inflation component, both, or neither. We reported component coefficients together with marginal predicted counts and zero probabilities.')
add_body('The final interpretation should distinguish four statements: evidence that the observed outcome is better represented by a zero-inflated model; evidence that contextual moderation improves the model; evidence that the moderation is localized to the count component, inflation component, or both; and the substantive size of the resulting change in predicted outcomes. These statements are related but are not interchangeable.')

add_heading('9 Specification checklist', level=1)
add_bullets([
    'Define the substantive meaning of the structural-zero state before fitting the model.',
    'Center and decompose the focal predictor so that within- and between-cluster associations are separated.',
    'Center the contextual moderator at a reported reference value and form the interaction from the within-cluster predictor.',
    'Specify count and inflation fixed effects separately and justify any difference between the two covariate sets.',
    'Choose the random-effects structure before testing focal moderation; consider inflation random intercepts when cluster zero heterogeneity is plausible.',
    'Use the same sample, offset, controls, and random-effects structure for nested component comparisons.',
    'Separate omnibus evidence for moderation from conditional evidence about component localization.',
    'Report valid-fit denominators and numerical failures separately from valid nonsignificant results.',
    'Translate both components into marginal expected counts and total-zero probabilities at meaningful covariate profiles.',
    'State whether inference is model-based, cluster-robust, HAC, or a sensitivity comparison, and report the number of clusters.',
])

# Final metadata and core properties
props = doc.core_properties
props.title = 'Model Specification of Contextual Moderation in Multilevel Zero Inflated Negative Binomial Models'
props.subject = 'Manuscript methods section for ML-ZINB contextual moderation'
props.author = 'OpenAI'
props.keywords = 'ML-ZINB, contextual moderation, multilevel models, zero inflation, model specification'

# Improve widow/orphan control and keep headings with following paragraphs
for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

# Save
doc.save(out)
print(out)
